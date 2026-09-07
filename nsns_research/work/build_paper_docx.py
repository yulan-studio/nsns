from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path

OUT = Path("outputs/运动训练对自闭症儿童焦虑及相关心理行为结局的影响_论文初稿.docx")
BLUE = "2E5E78"
LIGHT = "EAF1F5"
GRAY = "666666"


def set_font(run, size=10.5, bold=False, color="000000", latin="Arial", east="宋体"):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, 9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_font(run, 9, color=GRAY)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_reference(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-18)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(f"[{n}] {text}")
    set_font(r, 9.5)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(10.5)
for name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 6), ("Heading 3", 12, 8, 4)):
    st = styles[name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(BLUE)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.text = "研究论文初稿 | 结构化叙述综述"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    set_font(run, 8.5, color=GRAY)
add_page_number(sec.footer.paragraphs[0])

# Editorial-style manuscript title page.
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
r = p.add_run("结构化叙述综述 · 论文初稿")
set_font(r, 11, bold=True, color=BLUE, east="黑体")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
r = p.add_run("运动训练对自闭症儿童焦虑及相关心理行为结局的影响")
set_font(r, 24, bold=True, color="173B4D", east="黑体")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(26)
r = p.add_run("基于6项可获得全文研究的结构化叙述综述")
set_font(r, 14, color=BLUE, east="黑体")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
r = p.add_run("作者：____________________")
set_font(r, 11)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("单位：____________________    日期：2026年9月2日")
set_font(r, 10, color=GRAY)
doc.add_page_break()

add_heading(doc, "摘要", 1)
add_body(doc, "目的：总结结构化运动训练对自闭症儿童焦虑及相关心理行为结局的现有直接证据，并评价研究设计、缺失数据和安全性报告。", "目的：")
add_body(doc, "方法：基于既往PubMed检索和综述参考文献追溯形成的候选研究集，对能够合法获得全文的6项儿童运动干预研究进行结构化提取。提取研究对象、设计、运动剂量、对照、量表、效应结果、缺失数据和不良事件。随机试验参照RoB 2、非随机研究参照ROBINS-I的思路进行第一位评审者初评。由于检索与全文获取未达到完整系统综述要求，本研究采用分层叙述性综合，不计算合并效应。", "方法：")
add_body(doc, "结果：6项研究包括4项随机或随机主动对照试验、1项非随机对照研究和1项无对照单组研究。干预持续8至48周，涵盖有氧运动、蹦床、瑜伽、水中训练、学校运动以及协调、力量和平衡训练。两项采用SCARED的随机研究均报告运动组焦虑改善：Dhingra等研究干预后组间p=0.004；Luo等研究的SCARED变化中位数约为−9与0，FDR校正p<0.001，Cliff's delta=−0.92。Carey等无对照研究中教师与家长报告不一致。Ju等报告瑜伽后ABC问题行为总分下降，但总体偏倚风险较高。Marzouki等研究中情绪调节和情绪不稳定/负性的组别与时间交互均不显著。Toscano等非随机研究报告相关行为条目改善，但存在严重选择偏倚。除Luo等明确报告主动监测且两组均无不良事件外，其余研究的安全性报告普遍不充分。", "结果：")
add_body(doc, "结论：结构化运动可能改善部分自闭症儿童的焦虑及相关行为结局，但证据数量有限、结局构念高度异质，且若干研究存在较高偏倚风险。现阶段不能确定哪种运动方案最有效，也不能将结果用于替代个体化临床评估或直接决定治疗。", "结论：")
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)
r = p.add_run("关键词：")
set_font(r, bold=True)
r = p.add_run("自闭症；神经多样性；儿童；运动训练；焦虑；情绪调节；叙述综述")
set_font(r)

add_heading(doc, "1 引言", 1)
add_body(doc, "自闭症是一种神经发育差异，儿童在沟通、感官体验、行为调节和日常参与方面可能呈现多样化的支持需要。本研究采用神经多样性友好的表述，不把自闭症默认描述为必须“治愈”的缺陷。运动训练在儿童支持服务中具有参与性强、形式多样和可融入学校或社区环境等特点，其潜在作用可能涉及情绪调节、睡眠、日常功能和社会参与。")
add_body(doc, "焦虑及相关心理行为困难会影响部分自闭症儿童的学习、家庭生活和社会参与。已有运动研究采用SCARED、ASC-ASD、ABC、ERC和ATA等不同工具，但这些工具分别测量焦虑、问题行为、情绪调节或较广泛的自闭症相关特征，不能简单视为同一结局。此外，运动类型、对照条件和研究质量差异显著，容易造成过度概括。")
add_body(doc, "本研究旨在对6项已取得全文的儿童运动干预研究进行结构化叙述综合，重点回答三个问题：运动是否与焦虑或相关心理行为结局改善有关；现有结果在多大程度上受到研究偏倚和缺失数据影响；安全性信息是否足以支持稳健判断。本稿服务于临床研究论文写作，不用于诊断，也不为个体儿童直接决定干预方案。")

add_heading(doc, "2 方法", 1)
add_heading(doc, "2.1 研究定位", 2)
add_body(doc, "本稿定位为基于可获得全文证据的结构化叙述综述，而非完整系统综述或Meta分析。该定位源于7项潜在相关研究的全文未继续获取，且全部筛选记录尚未由第二位评审者独立复核。因此，研究不声称穷尽全部合格证据。")
add_heading(doc, "2.2 文献来源与筛选", 2)
add_body(doc, "前期通过PubMed可复现检索以及3篇开放全文综述的参考文献追溯识别候选研究。参考文献追溯产生57条记录，去重后为53篇，并补入数据库检索发现的2篇研究，共形成55篇候选记录。第一位评审者暂定13篇为主要或相关精神健康证据，另有3篇睡眠次要结局研究。本稿仅纳入其中已取得并核对全文的6篇主要或相关研究；其余7篇因本阶段决定不继续获取全文而未进入综合，3篇睡眠研究不属于本稿主要分析范围。PubMed初次检索的303条记录与参考文献来源尚未完成统一去重，因此不将二者直接相加。")
add_heading(doc, "2.3 纳入标准", 2)
add_body(doc, "研究对象为3至18岁自闭症儿童或青少年；干预为至少4周的结构化、重复运动训练；报告焦虑、抑郁、情绪调节、内化/外化行为、心理困扰、幸福感或心理相关生活质量等结局；研究为原始干预研究。运动与认知行为治疗、心理教育或其他复杂治疗无法分离的多组分研究不纳入主要综合。")
add_heading(doc, "2.4 数据提取与证据分层", 2)
add_body(doc, "逐篇提取样本、年龄、研究设计、运动类型与剂量、对照条件、结局工具、结果数值、效应量、置信区间、P值、缺失数据和不良事件。证据分为核心焦虑证据、相关精神健康或行为证据以及探索性证据。只有原始全文明确报告的数值才作为直接证据；论文未报告的信息标记为“未报告”，不得推断为0。")
add_heading(doc, "2.5 偏倚风险与综合方法", 2)
add_body(doc, "随机试验参考RoB 2框架，从随机过程、偏离预期干预、缺失结局、结局测量和选择性报告方面进行初评；非随机研究参考ROBINS-I思路关注混杂、选择、偏离干预、缺失、测量与报告。该评价由第一位评审者完成，属于初步判断。鉴于结局工具和研究设计高度异质，本研究按焦虑、情绪调节和问题行为分层叙述，不计算合并效应。")

add_heading(doc, "3 结果", 1)
add_heading(doc, "3.1 研究选择", 2)
add_body(doc, "本阶段综合6项已取得全文的研究[1-6]。另有7项潜在相关研究未获取全文，因此未纳入当前结果；3项睡眠研究被预先划为次要结局，不属于本稿主要分析。该选择方式可能造成全文获取偏倚，结果应解释为可获得全文样本的证据概貌。")
add_heading(doc, "3.2 研究特征", 2)
add_body(doc, "6项研究发表于2022至2026年，地点包括印度、爱尔兰、中国、突尼斯和巴西。样本规模从17至229名不等；干预持续8至48周。运动类型和测量工具差异明显。表1概述研究特征。")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("表1  纳入研究的主要特征")
set_font(r, 10, bold=True)
rows = [
    ["研究","设计/样本","运动方案","心理行为结局与主要发现"],
    ["Dhingra 2025[1]","RCT；n=38；7–13岁","有氧；8周；3次/周；30分钟","SCARED；干预后组间p=0.004"],
    ["Carey 2022[2]","单组；入组47；完整分析各20名","学校运动；16周；3次/周；60分钟","ASC-ASD；教师改善、家长不显著"],
    ["Ju 2024[3]","RCT；n=17","瑜伽；8周；3次/周；45–50分钟","ABC；8周总分77.89±10.57 vs 95.88±7.77"],
    ["Marzouki 2022[4]","三组RCT；随机28；分析22","水中训练；8周；2次/周；50分钟","ERC；两个组别×时间交互均不显著"],
    ["Toscano 2022[5]","非随机对照；n=229","协调/力量/平衡；48周；2次/周","ATA；反应性等改善，但选择偏倚严重"],
    ["Luo 2026[6]","随机主动对照；n=50；9–14岁","蹦床；8周；3次/周；30分钟","SCARED；变化约−9 vs 0；delta=−0.92"]
]
table = doc.add_table(rows=len(rows), cols=4)
table.style = "Table Grid"
tr_pr = table.rows[0]._tr.get_or_add_trPr()
tbl_header = OxmlElement("w:tblHeader")
tbl_header.set(qn("w:val"), "true")
tr_pr.append(tbl_header)
for i, row in enumerate(rows):
    for j, value in enumerate(row):
        cell = table.cell(i, j)
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if i == 0:
            shade(cell, LIGHT)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                set_font(run, 8.2, bold=(i == 0), east="宋体")
set_table_geometry(table, [1450, 2100, 2250, 3560])

add_heading(doc, "3.3 焦虑结局", 2)
add_body(doc, "Dhingra等[1]将38名7至13岁儿童随机分为有氧训练组与知识讲座对照组。运动组接受8周、每周3次、每次30分钟训练，目标强度为年龄推算最大心率的65%至85%。运动组SCARED中位数由44（四分位距42–50）降至42（41–46），对照组由47（45–52）变为51（44–53）；干预后组间p=0.004。研究没有提供组间效应量及95%置信区间，限制了与其他试验的定量比较。38名参与者均完成研究。研究说明使用训练日记记录不良事件，但结果没有公布事件数量，故本综述记为“未报告”，而非0例。")
add_body(doc, "Luo等[6]将50名9至14岁儿童随机分至蹦床运动和弹力带主动对照组。两组均训练8周、每周3次、每次30分钟。两组SCARED变化中位数约为−9和0，FDR校正p<0.001，Cliff's delta=−0.92。50名参与者全部完成，依从性约为85%和82%；研究主动监测并报告两组均无不良事件。该研究是本阶段报告最完整的直接焦虑证据，但仍受单中心和小样本限制。")
add_body(doc, "Carey等[2]开展16周学校运动单组研究。47名儿童入组，16周后只有24名至少具有家长或教师资料，家长和教师的三时间点完整分析各为20名。教师报告的焦虑有所改善，家长报告未见显著变化。由于缺乏对照组、报告者结果不一致且结局流失率较高，观察到的时间变化不能被解释为确定的运动因果效应。")

add_heading(doc, "3.4 情绪调节和问题行为", 2)
add_body(doc, "Ju等[3]在17名儿童中比较8周团体瑜伽和日常课程活动。8周后ABC总分为77.89±10.57和95.88±7.77，Bonferroni校正p=0.046，论文报告Cohen's d=1.847；易激惹和社会退缩亦呈组间差异。由于样本极小，且教师知晓分组并评定主观结局，该结果存在较高偏倚风险。ABC测量的问题行为不能直接等同于焦虑或抑郁。")
add_body(doc, "Marzouki等[4]将28名6至7岁儿童随机分为技术型水中训练、游戏型水中训练和常规活动对照组，最终分析22名。情绪调节的组别与时间交互不显著（F(2,19)=0.826，p=0.453），情绪不稳定/负性的交互亦不显著（F(2,19)=1.641，p=0.220）。因此，该研究没有证明水中训练在情绪调节方面优于对照。6名参与者因研究期间出现的健康问题被排除，作者称这些问题并非训练造成，但未提供完整的分组事件表。")
add_body(doc, "Toscano等[5]纳入229名2.3至17.3岁参与者，运动组127名、同机构对照62名、另一机构对照40名。干预包括协调、力量和平衡训练，持续48周。作者使用贝叶斯多层模型并报告运动组反应性等ATA条目改善。然而，家庭按参与意愿选择组别，且仅将出席率达到90%者纳入分析，选择偏倚和残余混杂较严重。该研究只能作为支持性关联证据。")

add_heading(doc, "3.5 缺失数据、安全性与偏倚风险", 2)
add_body(doc, "缺失数据处理差异明显。Dhingra和Luo研究报告全部参与者完成；Marzouki研究随机后排除6/28名参与者；Carey研究仅对完整案例进行分析；Toscano研究只纳入达到出席率阈值者。安全性方面，只有Luo研究清楚报告主动监测及两组0例事件。其他研究未报告或报告不完整，不能据此认定运动没有风险。")
add_body(doc, "第一位评审者初评认为，Luo和Dhingra研究总体存在“有些担忧”；Ju和Marzouki随机试验总体偏倚风险较高；Carey和Toscano非随机研究存在严重偏倚风险。该判断需要第二位评审者按照完整信号问题独立复核。")

add_heading(doc, "4 讨论", 1)
add_heading(doc, "4.1 主要发现", 2)
add_body(doc, "本阶段最直接的证据来自两项采用SCARED的随机研究，两者均提示运动后焦虑改善。然而，两项研究的对照条件不同：一项使用知识讲座，另一项使用弹力带主动训练；Dhingra研究又缺少可直接合并的组间效应量和95%置信区间。因此，方向一致并不等于能够得到可靠的合并效应。")
add_body(doc, "相关行为结局的结果并不完全一致。Ju研究提示瑜伽可能改善易激惹和社会退缩，但偏倚风险较高；Marzouki研究没有观察到水中训练相对对照改善情绪调节；Toscano研究样本较大，却因非随机分组而难以排除家庭动机、资源和基线差异等混杂。这些结果强调，不能把所有心理和行为量表合并成一个笼统的“精神健康”结局。")
add_heading(doc, "4.2 临床与研究意义", 2)
add_body(doc, "运动训练可被理解为支持儿童参与、体能、情绪体验和日常生活的一种可选方案，而不是纠正神经多样性的工具。现有证据不支持宣称某一种运动适用于所有自闭症儿童。实践中应尊重儿童偏好、感官需求、沟通方式和共同决策，并结合专业人员对运动安全、环境适配和可持续参与的评估。")
add_body(doc, "未来研究应预注册主要结局，清楚报告随机与分配隐藏，使用盲法或尽可能独立的结局评定者，报告每个时间点的样本和缺失原因，并系统记录不良事件。研究还应分别报告儿童自评、家长评定、教师评定和生理指标，避免把不同信息来源混为同一结果。")
add_heading(doc, "4.3 为什么暂不进行Meta分析", 2)
add_body(doc, "当前只有两项研究使用SCARED，且Dhingra研究的效应信息不完整；Carey研究没有对照组；ABC、ERC和ATA又分别测量问题行为、情绪调节和较广泛的自闭症相关特征。研究设计、对照和偏倚风险亦不一致。强行合并会产生难以解释的数字精确性，因此本稿保留分层叙述性综合。")

add_heading(doc, "5 局限性", 1)
add_body(doc, "本研究最重要的局限是全文获取不完整。13项主要或相关候选研究中只有6项进入本稿，另7项未继续获取全文。这可能造成全文获取偏倚，使结果不能代表全部相关研究。初始PubMed记录与参考文献追溯记录尚未完成统一去重；筛选和偏倚风险评价也尚未由第二位评审者独立完成。因此，本稿不能称为完整系统综述，也不能绘制声称完整检索流程的正式PRISMA纳入图。")
add_body(doc, "纳入研究本身也存在局限，包括样本小、量表构念不一致、部分研究缺乏对照、随机后排除参与者、非随机分组、主观结局评定和安全性报告不足。本稿没有纳入睡眠次要结局，也没有比较不同支持需要等级、年龄、性别或运动剂量的亚组效果。")

add_heading(doc, "6 结论", 1)
add_body(doc, "基于6项可获得全文研究，结构化运动可能对部分自闭症儿童的焦虑及相关问题行为有益，但情绪调节结果并不一致，整体证据确定性有限。当前资料不足以确定最佳运动类型、剂量或长期效果，也不足以计算可信的总体效应。后续若扩大为完整系统综述，应补齐其余候选全文、完成双人独立筛选与偏倚评价，并按结局构念分别综合。")

add_heading(doc, "声明", 1)
add_body(doc, "研究用途：本稿用于临床研究与学术写作，不构成诊断、治疗建议或个体化医疗决策。数据边界：本稿仅使用公开论文中的汇总结果，不处理可识别儿童身份的信息。证据表达：直接证据、缺失报告和评审者推论已尽量分开；未报告的数值或事件不被推定为0。利益冲突与基金信息由作者投稿前补充。")

add_heading(doc, "参考文献", 1)
add_reference(doc, 1, "Dhingra A, Gupta A, Aditi. Effect of Aerobic Training on Anxiety, Activities of Daily Living and Visual Motor Skills Performance in Children with Autism: A Randomised Controlled Trial. Journal of Clinical and Diagnostic Research. 2025;19(2):SC06-SC11. doi:10.7860/JCDR/2025/73346.20624.")
add_reference(doc, 2, "Carey M, Sheehan D, Healy S, Knott F, Kinsella S. The Effects of a 16-Week School-Based Exercise Program on Anxiety in Children with Autism Spectrum Disorder. International Journal of Environmental Research and Public Health. 2022;19(9):5471. PMID:35564866. doi:10.3390/ijerph19095471.")
add_reference(doc, 3, "Ju X, Liu H, Xu J, Hu B, Jin Y, Lu C. Effect of Yoga Intervention on Problem Behavior and Motor Coordination in Children with Autism. Behavioral Sciences. 2024;14(2):116. PMID:38392469. doi:10.3390/bs14020116.")
add_reference(doc, 4, "Marzouki H, Soussi B, Selmi O, Hajji Y, Marsigliante S, Bouhlel E, et al. Effects of Aquatic Training in Children with Autism Spectrum Disorder. Biology. 2022;11(5):657. PMID:35625385. doi:10.3390/biology11050657.")
add_reference(doc, 5, "Toscano CVA, Ferreira JP, Quinaud RT, Silva KMN, Carvalho HM, Gaspar JM. Exercise improves the social and behavioral skills of children and adolescent with autism spectrum disorders. Frontiers in Psychiatry. 2022;13:1027799. PMID:36620673. doi:10.3389/fpsyt.2022.1027799.")
add_reference(doc, 6, "Luo B, Wanpen S, Eungpinichpong W. Trampoline exercise reduces anxiety and improves motor skills in children with autism spectrum disorder: A randomized controlled trial. Medicine. 2026;105(19):e48616. PMID:42116390. doi:10.1097/MD.0000000000048616.")

doc.core_properties.title = "运动训练对自闭症儿童焦虑及相关心理行为结局的影响"
doc.core_properties.subject = "基于6项可获得全文研究的结构化叙述综述"
doc.core_properties.author = ""
doc.core_properties.keywords = "自闭症, 神经多样性, 儿童, 运动训练, 焦虑"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
